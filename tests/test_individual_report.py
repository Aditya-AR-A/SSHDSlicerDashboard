"""
Unit tests for Individual Performance Report generator.
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
import pandas as pd
import openpyxl

from slicing_dashboard.reporting.individual import generate_individual_report_data
from slicing_dashboard.reporting.excel import write_individual_report


class TestIndividualReport(unittest.TestCase):
    def setUp(self):
        master_records = [
            # Normal completed task
            {
                "id": "task_1",
                "video_id": "vid_1",
                "slice_batch": 101,
                "status": "slice_completed",
                "duration_seconds": 120.0,
                "completed_date": "2026-09-01",
                "completed_at": "2026-09-01 10:00:00",
                "user_name": "Priya",
                "user_id": "SSHD-Priya",
                "is_completed": True,
                "completion_type": "normal",
                "rework_by": None,
            },
            # Error completed task
            {
                "id": "task_2",
                "video_id": "vid_2",
                "slice_batch": 101,
                "status": "video_error_confirmed",
                "duration_seconds": 80.0,
                "completed_date": "2026-09-01",
                "completed_at": "2026-09-01 11:00:00",
                "user_name": "Priya",
                "user_id": "SSHD-Slicer2",
                "is_completed": True,
                "completion_type": "error",
                "rework_by": None,
            },
            # In-progress assigned task (partial batch)
            {
                "id": "task_3",
                "video_id": "vid_3",
                "slice_batch": 101,
                "status": "slicing",
                "duration_seconds": 90.0,
                "completed_date": "2026-09-01",
                "completed_at": None,
                "user_name": "Priya",
                "user_id": "SSHD-Priya",
                "is_completed": False,
                "completion_type": "none",
                "rework_by": None,
            },
            # Task for another user
            {
                "id": "task_other",
                "video_id": "vid_other",
                "slice_batch": 102,
                "status": "slice_completed",
                "duration_seconds": 300.0,
                "completed_date": "2026-09-01",
                "completed_at": "2026-09-01 12:00:00",
                "user_name": "Aditya",
                "user_id": "SSHD-Aditya",
                "is_completed": True,
                "completion_type": "normal",
                "rework_by": None,
            },
        ]
        self.master_df = pd.DataFrame(master_records)

        transition_records = [
            {"task_id": "task_1", "user_name": "SSHD-Priya", "duration_seconds": 120.0, "date": "2026-09-01", "type": "submitted"},
            {"task_id": "task_1", "user_name": "SSHD-Priya", "duration_seconds": 120.0, "date": "2026-09-01", "type": "leader_passed"},
            {"task_id": "task_2", "user_name": "SSHD-Slicer2", "duration_seconds": 80.0, "date": "2026-08-30", "type": "submitted"},
            {"task_id": "task_2", "user_name": "SSHD-Slicer2", "duration_seconds": 80.0, "date": "2026-08-31", "type": "leader_returned"},
            {"task_id": "task_2", "user_name": "SSHD-Slicer2", "duration_seconds": 80.0, "date": "2026-09-01", "type": "submitted"},
        ]
        self.transitions_df = pd.DataFrame(transition_records)

        self.user_mapping = {
            "SSHD-Priya": "Priya",
            "SSHD-Slicer2": "Priya",
            "SSHD-Aditya": "Aditya",
        }

    def test_generate_individual_report_data(self):
        report_data = generate_individual_report_data(
            master_df=self.master_df,
            transitions_df=self.transitions_df,
            individual="Priya",
            target_date="2026-09-01",
            target_month="2026-09",
            user_mapping=self.user_mapping,
        )

        summary = report_data["summary"]
        self.assertEqual(summary["individual"], "Priya")
        self.assertEqual(summary["total_completed_count"], 2)
        self.assertEqual(summary["completed_normal_count"], 1)
        self.assertEqual(summary["completed_error_count"], 1)
        self.assertEqual(summary["total_assigned_count"], 1)
        self.assertEqual(summary["total_completed_duration_seconds"], 200.0)

        # Monthly Daily Distribution covers all 30 days of September
        monthly_daily = report_data["monthly_daily"]
        self.assertEqual(len(monthly_daily), 30)
        self.assertIn("Date", monthly_daily.columns)
        self.assertIn("Pending Leader Hours", monthly_daily.columns)
        self.assertIn("Total Completed Hours", monthly_daily.columns)

        # Check day 2026-09-01
        day_row = monthly_daily[monthly_daily["Date"] == "2026-09-01"].iloc[0]
        self.assertEqual(day_row["Total Completed Count"], 2)
        self.assertEqual(day_row["Completed Normal Count"], 1)
        self.assertEqual(day_row["Completed Error Count"], 1)

        # Today's Activity checks
        ts = report_data["today_summary"]
        self.assertEqual(ts["completed_today_count"], 2)
        self.assertEqual(ts["assigned_today_count"], 1)
        self.assertEqual(ts["first_time_count"], 2)
        self.assertEqual(ts["rework_count"], 1)

        # Video Detail checks
        vd = report_data["today_video_detail"]
        self.assertEqual(len(vd), 3)
        t2_detail = vd[vd["Task ID"] == "task_2"].iloc[0]
        self.assertEqual(t2_detail["Is Rework"], "Yes")
        self.assertEqual(t2_detail["Returned By"], "Leader")

    def test_write_individual_report(self):
        report_data = generate_individual_report_data(
            master_df=self.master_df,
            transitions_df=self.transitions_df,
            individual="Priya",
            target_date="2026-09-01",
            target_month="2026-09",
            user_mapping=self.user_mapping,
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = write_individual_report(report_data, output_dir=tmpdir)
            self.assertTrue(Path(out_path).exists())

            wb = openpyxl.load_workbook(out_path)
            expected_sheets = [
                "Individual Summary",
                "Monthly Daily Distribution",
                "Today Activity",
                "Today Video Detail",
                "Data Helper",
            ]
            self.assertEqual(wb.sheetnames, expected_sheets)

            ws_summary = wb["Individual Summary"]
            self.assertEqual(ws_summary["A1"].value, "Individual Performance Summary: Priya")

            ws_daily = wb["Monthly Daily Distribution"]
            tot_row = ws_daily.max_row
            self.assertEqual(ws_daily.cell(row=tot_row, column=1).value, "Total")
            self.assertIn("=SUM(", str(ws_daily.cell(row=tot_row, column=2).value))


if __name__ == "__main__":
    unittest.main()
